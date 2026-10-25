"""Generate the Hostel Facilities Request Letter as an editable .docx file.

Run with the project's virtual environment:
    ./.venv/bin/python generate_hostel_letter.py

This produces Hostel_Facilities_Request_Letter.docx in the same directory.
The script is kept in the repo so the document stays regenerable/editable.
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Cm

OUTPUT_FILENAME = "Hostel_Facilities_Request_Letter.docx"

RECIPIENTS = [
    "Dr. R. S. Bhati Sir",
    "Mr. Anil Srivastava Sir",
    "Mr. Pankaj Kushwaha Sir",
]

ISSUES = [
    (
        "Additional Refrigerator",
        "The existing refrigerator facility is insufficient for the number of "
        "students staying in the hostel. Also, the ground-floor refrigerator is "
        "being used by the family residing downstairs, which creates "
        "inconvenience for hostel students. We request an additional "
        "refrigerator and proper arrangements to ensure that students have "
        "adequate access to this facility.",
    ),
    (
        "Additional Washing Machine",
        "With so many students sharing the existing washing machine, it becomes "
        "difficult for everyone to manage their laundry conveniently. We request "
        "an additional washing machine as soon as possible.",
    ),
    (
        "Pending Repair Work",
        "Furniture repairs, plumbing work, and electrical repairs are still not "
        "being done properly despite repeated requests and reminders to the "
        "authorities. We request that these issues be addressed and the "
        "necessary repairs completed without further delay.",
    ),
]


def set_cell_borders(cell):
    """Add single-line borders to all four sides of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    tc_pr.append(borders)


def set_row_height(row, cm_value):
    """Set a minimum row height so students have room to write and sign."""
    tr_pr = row._tr.get_or_add_trPr()
    tr_height = OxmlElement("w:trHeight")
    # 1 cm = 567 twips
    tr_height.set(qn("w:val"), str(int(cm_value * 567)))
    tr_height.set(qn("w:hRule"), "atLeast")
    tr_pr.append(tr_height)


def build_document():
    doc = Document()

    # Base font: simple and readable.
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # --- Date ---
    doc.add_paragraph("Date: __________")
    doc.add_paragraph()

    # --- Recipients ---
    doc.add_paragraph("To,")
    for name in RECIPIENTS:
        doc.add_paragraph(name)
    doc.add_paragraph()

    # --- Subject ---
    subject = doc.add_paragraph()
    run = subject.add_run(
        "Subject: Urgent Request for Hostel Facilities and Pending Repairs"
    )
    run.bold = True
    doc.add_paragraph()

    # --- Salutation ---
    doc.add_paragraph("Respected Sirs,")
    doc.add_paragraph()

    # --- Opening line ---
    doc.add_paragraph(
        "We, the students residing in Hostel 2, Patrakar Colony, Man Hostel, "
        "would like to bring a few important hostel-related issues to your "
        "attention."
    )
    doc.add_paragraph()

    # --- Numbered issues with bold sub-headings ---
    for index, (heading, body) in enumerate(ISSUES, start=1):
        paragraph = doc.add_paragraph()
        heading_run = paragraph.add_run(f"{index}. {heading}: ")
        heading_run.bold = True
        paragraph.add_run(body)
        doc.add_paragraph()

    # --- Closing paragraphs ---
    doc.add_paragraph(
        "These are basic facilities and maintenance requirements that directly "
        "affect our daily lives in the hostel. We request you to kindly look "
        "into these matters and take the necessary action at the earliest."
    )
    doc.add_paragraph()
    doc.add_paragraph("Thank you for your understanding and support.")
    doc.add_paragraph()
    doc.add_paragraph("Yours sincerely,")
    doc.add_paragraph(
        "Students of Hostel 2, Patrakar Colony, Man Hostel"
    )

    # --- Page break before the signature section ---
    doc.add_page_break()

    # --- Signature section heading ---
    heading_para = doc.add_paragraph()
    heading_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading_run = heading_para.add_run("STUDENT NAMES AND SIGNATURES")
    heading_run.bold = True
    heading_run.font.size = Pt(13)

    subheading_para = doc.add_paragraph()
    subheading_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subheading_run = subheading_para.add_run(
        "Hostel 2, Patrakar Colony, Man Hostel"
    )
    subheading_run.bold = True
    doc.add_paragraph()

    # --- Signature table: header + 15 numbered rows ---
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    header_cells = table.rows[0].cells
    headers = ["SR. NO.", "NAME OF STUDENT", "SIGNATURE"]
    for cell, text in zip(header_cells, headers):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_borders(cell)

    for number in range(1, 16):
        row = table.add_row()
        cells = row.cells
        cells[0].text = str(number)
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        # Name and signature cells stay blank for handwriting / image insert.
        cells[1].text = ""
        cells[2].text = ""
        for cell in cells:
            set_cell_borders(cell)
        # Generous height so students have room to write and sign.
        set_row_height(row, 1.1)

    # Column widths: narrow SR. NO., wider name and signature columns.
    widths = [Cm(2.0), Cm(8.0), Cm(6.0)]
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width

    doc.save(OUTPUT_FILENAME)
    print(f"Created {OUTPUT_FILENAME}")


if __name__ == "__main__":
    build_document()
